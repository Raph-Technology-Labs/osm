# backend/app/routers/part_details.py
"""
Part Details screen -- browse / search / edit / export configured parts.

Mounted at /api/v1/part-details. Its own prefix rather than /parts or
/dashboard, because parts.py, parts_admin.py and dashboard.py already own
those and part_details registers LAST -- a colliding path would silently
lose to whichever router came first.

    GET    /part-details/parts?category_id=&pipeline=
    GET    /part-details/parts/{part_id}
    GET    /part-details/by-code?part_code=
    GET    /part-details/parts/{part_id}/image
    PUT    /part-details/parts/{part_id}              [admin]
    DELETE /part-details/parts/{part_id}              [admin]
    GET    /part-details/export

Categories come from the existing GET /parts/categories -- not duplicated here.

Image: Part.image is LargeBinary and is never inlined in a listing (the
PartOut comment in schemas.py). Rows carry has_image; the browser pulls and
caches each thumbnail. Uploads arrive as a base64 data URI in the PUT body.
"""

from __future__ import annotations

import io
from datetime import datetime
from typing import List, Optional

from fastapi import APIRouter, Depends, HTTPException, Query, Response, status
from fastapi.responses import StreamingResponse
from sqlalchemy import func
from sqlalchemy.orm import Session, joinedload

from app.auth.dependencies import require_role
from app.db.db import SessionLocal
from app.models.models import Category, Part, PartConfig, PartSession
from app.schemas import PartDetailOut, PartUpdate


router = APIRouter(prefix="/part-details", tags=["part-details"])

PIPELINE_FILTERS = ("defect", "measurement", "both", "none")


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


# ─────────────────────────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────────────────────────

def _row(part: Part) -> dict:
    """Part ORM object -> the flat shape the table renders."""
    return {
        "part_id": part.part_id,
        "part_code": part.part_code,
        "part_name": part.part_name,
        "category_id": part.category_id,
        "category_name": part.category.category_name if part.category else None,
        "part_weight": part.part_weight,
        "notes": part.notes,
        "dimensions": part.dimensions or {},
        "defects": part.defects or {},
        "has_image": part.image is not None,
        "has_defect_pipeline": part.has_defect_pipeline,
        "has_measurement_pipeline": part.has_measurement_pipeline,
        "active_config_version": (
            part.active_config.version if part.active_config else None
        ),
        "created_at": part.created_at,
        "created_by_name": part.creator.name if part.creator else None,
    }


def _base_query(db: Session):
    return db.query(Part).options(
        joinedload(Part.category),
        joinedload(Part.creator),
        joinedload(Part.active_config),
    )


def _get_part_or_404(db: Session, part_id: int) -> Part:
    part = _base_query(db).filter(Part.part_id == part_id).first()
    if part is None:
        raise HTTPException(
            status.HTTP_404_NOT_FOUND,
            detail="Part not found",
        )
    return part


def _image_media_type(blob: bytes) -> str:
    return "image/png" if blob.startswith(b"\x89PNG") else "image/jpeg"


# ─────────────────────────────────────────────────────────────────
# Listing / lookup
# ─────────────────────────────────────────────────────────────────

@router.get("/parts", response_model=List[PartDetailOut])
def list_parts(
    category_id: Optional[int] = Query(None),
    pipeline: Optional[str] = Query(
        None,
        description="defect | measurement | both | none",
    ),
    db: Session = Depends(get_db),
):
    q = _base_query(db)

    if category_id is not None:
        q = q.filter(Part.category_id == category_id)

    if pipeline:
        if pipeline not in PIPELINE_FILTERS:
            raise HTTPException(
                status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"pipeline must be one of {PIPELINE_FILTERS}",
            )

        if pipeline == "defect":
            q = q.filter(Part.has_defect_pipeline.is_(True))

        elif pipeline == "measurement":
            q = q.filter(Part.has_measurement_pipeline.is_(True))

        elif pipeline == "both":
            q = q.filter(
                Part.has_defect_pipeline.is_(True),
                Part.has_measurement_pipeline.is_(True),
            )

        elif pipeline == "none":
            q = q.filter(
                Part.has_defect_pipeline.is_(False),
                Part.has_measurement_pipeline.is_(False),
            )

    return [_row(p) for p in q.order_by(Part.part_code).all()]


@router.get("/by-code", response_model=PartDetailOut)
def part_by_code(
    part_code: str = Query(..., min_length=1),
    db: Session = Depends(get_db),
):
    """Barcode scan lookup -- case-insensitive exact match."""
    part = (
        _base_query(db)
        .filter(func.lower(Part.part_code) == part_code.strip().lower())
        .first()
    )

    if part is None:
        raise HTTPException(
            status.HTTP_404_NOT_FOUND,
            detail=f"No part with code '{part_code}'",
        )

    return _row(part)


@router.get("/parts/{part_id}/image")
def get_part_image(
    part_id: int,
    db: Session = Depends(get_db),
):
    part = db.query(Part).filter(Part.part_id == part_id).first()

    if part is None or part.image is None:
        # 204 keeps the <img> quiet instead of logging a 404 per row
        return Response(status_code=status.HTTP_204_NO_CONTENT)

    return Response(
        content=part.image,
        media_type=_image_media_type(part.image),
        headers={"Cache-Control": "private, max-age=300"},
    )


@router.get("/parts/{part_id}", response_model=PartDetailOut)
def get_part(
    part_id: int,
    db: Session = Depends(get_db),
):
    return _row(_get_part_or_404(db, part_id))


# ─────────────────────────────────────────────────────────────────
# Update [administrator | superadministrator]
# ─────────────────────────────────────────────────────────────────

@router.put("/parts/{part_id}", response_model=PartDetailOut)
def update_part(
    part_id: int,
    payload: PartUpdate,
    db: Session = Depends(get_db),
    _admin=Depends(require_role("administrator")),
):
    part = _get_part_or_404(db, part_id)

    if payload.category_id is not None:
        exists = (
            db.query(Category.category_id)
            .filter(Category.category_id == payload.category_id)
            .first()
        )

        if not exists:
            raise HTTPException(
                status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Category {payload.category_id} does not exist",
            )

    try:
        new_image = payload.decoded_image()
    except ValueError as exc:
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=str(exc),
        )

    part.part_name = payload.part_name.strip()
    part.category_id = payload.category_id
    part.part_weight = payload.part_weight
    part.notes = payload.notes

    # exclude_none keeps the stored JSON in the same sparse shape
    # the Add Part form writes -- absent keys rather than explicit nulls.
    part.dimensions = {
        k: v.model_dump(exclude_none=True)
        for k, v in payload.dimensions.items()
    } or None

    part.defects = {
        k: v.model_dump(exclude_none=True)
        for k, v in payload.defects.items()
    } or None

    if payload.clear_image:
        part.image = None
    elif new_image is not None:
        part.image = new_image

    db.commit()
    db.refresh(part)

    return _row(part)


# ─────────────────────────────────────────────────────────────────
# Delete [administrator | superadministrator]
# ─────────────────────────────────────────────────────────────────

@router.delete("/parts/{part_id}")
def delete_part(
    part_id: int,
    db: Session = Depends(get_db),
    _admin=Depends(require_role("administrator")),
):
    part = _get_part_or_404(db, part_id)

    # Production history must not be orphaned --
    # PartSession.part_id is NOT NULL.
    session_count = (
        db.query(func.count(PartSession.id))
        .filter(PartSession.part_id == part_id)
        .scalar()
    )

    if session_count:
        raise HTTPException(
            status.HTTP_409_CONFLICT,
            detail=(
                f"Cannot delete '{part.part_code}' -- "
                f"{session_count} production session(s) reference it. "
                f"Archive the part instead."
            ),
        )

    code = part.part_code

    # Break the parts.active_config_id -> part_configs.id cycle first.
    part.active_config_id = None
    db.flush()

    db.query(PartConfig).filter(
        PartConfig.part_id == part_id
    ).delete(
        synchronize_session=False
    )

    db.delete(part)
    db.commit()

    return {"message": f"Part '{code}' deleted"}


# ─────────────────────────────────────────────────────────────────
# Excel export
# ─────────────────────────────────────────────────────────────────

@router.get("/export")
def export_parts(db: Session = Depends(get_db)):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HTTPException(
            status.HTTP_501_NOT_IMPLEMENTED,
            detail="openpyxl is not installed on the server",
        )

    parts = _base_query(db).order_by(Part.part_code).all()

    wb = Workbook()

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="37474F")

    def write_sheet(ws, headers, rows):
        ws.append(headers)

        for cell in ws[1]:
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center")

        for r in rows:
            ws.append(r)

        for i, h in enumerate(headers, start=1):
            width = (
                max(
                    [len(str(h))]
                    + [len(str(r[i - 1])) for r in rows]
                    or [0]
                )
                + 3
            )

            ws.column_dimensions[
                get_column_letter(i)
            ].width = min(width, 45)

        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Parts"

    write_sheet(
        ws,
        [
            "Part Code",
            "Part Name",
            "Category",
            "Weight (g)",
            "Defect Pipeline",
            "Measurement Pipeline",
            "Active Config v",
            "Params",
            "Defects",
            "Created By",
            "Created At",
            "Notes",
        ],
        [
            [
                p.part_code,
                p.part_name,
                p.category.category_name if p.category else "",
                p.part_weight if p.part_weight is not None else "",
                "Yes" if p.has_defect_pipeline else "No",
                "Yes" if p.has_measurement_pipeline else "No",
                p.active_config.version if p.active_config else "",
                len(p.dimensions or {}),
                len(p.defects or {}),
                p.creator.name if p.creator else "",
                (
                    p.created_at.strftime("%Y-%m-%d %H:%M")
                    if p.created_at
                    else ""
                ),
                p.notes or "",
            ]
            for p in parts
        ],
    )

    write_sheet(
        wb.create_sheet("Dimensions"),
        [
            "Part Code",
            "Parameter",
            "Nominal",
            "Lower Tolerance",
            "Upper Tolerance",
            "Unit",
            "Cal. Factor",
            "Notes",
        ],
        [
            [
                p.part_code,
                name,
                spec.get("nominal", ""),
                spec.get("lower_limit", ""),
                spec.get("upper_limit", ""),
                spec.get("unit", ""),
                spec.get("calibration_factor", ""),
                spec.get("notes", ""),
            ]
            for p in parts
            for name, spec in sorted(
                (p.dimensions or {}).items()
            )
        ],
    )

    write_sheet(
        wb.create_sheet("Defects"),
        [
            "Part Code",
            "Defect",
            "Confidence Threshold",
            "Severity",
            "Notes",
        ],
        [
            [
                p.part_code,
                name,
                spec.get("conf_thresh", ""),
                spec.get("severity", ""),
                spec.get("notes", ""),
            ]
            for p in parts
            for name, spec in sorted(
                (p.defects or {}).items()
            )
        ],
    )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"OSM_PartsData_{datetime.now():%Y-%m-%d}.xlsx"

    return StreamingResponse(
        buf,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"'
        },
    )